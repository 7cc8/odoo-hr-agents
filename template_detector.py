"""
Template Detector - كاشف نوع Template تلقائياً
==================================================
يفحص أي ملف Excel ويحدد:
- هل هو Master Template (multi-sheet)
- أم Modular Template (single-purpose)
- أي sheets موجودة
- أي وكلاء يحتاجون للمعالجة
"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Dict, Any, List, Optional, Set
from openpyxl import load_workbook

logger = logging.getLogger("template_detector")


# ═══════════════════════════════════════════════════════════════
# Sheet Type Mapping - ربط الـ sheets بالوكلاء
# ═══════════════════════════════════════════════════════════════

SHEET_TYPE_MAP = {
    # Company Setup
    "company_setup": {
        "patterns": ["company", "setup", "إعداد", "شركة", "منشأة"],
        "agents": ["fahas", "katib"],
        "priority": 1,  # أول شيء
    },
    
    # Employees (الأكثر شيوعاً)
    "employees": {
        "patterns": ["employee", "موظف", "موظفين", "staff", "personnel"],
        "agents": ["fahas", "musleh", "hares", "katib"],
        "priority": 2,
    },
    
    # Contracts
    "contracts": {
        "patterns": ["contract", "عقد", "عقود", "allowance", "بدل"],
        "agents": ["fahas", "mudaqqiq", "katib"],
        "priority": 3,
    },
    
    # Payroll & GOSI
    "payroll": {
        "patterns": ["payroll", "salary", "gosi", "رواتب", "مرتب", "تأمين"],
        "agents": ["fahas", "mudaqqiq", "katib"],
        "priority": 4,
    },
    
    # Time Off / Leave
    "timeoff": {
        "patterns": ["time", "off", "leave", "إجازة", "إجازات", "vacation"],
        "agents": ["fahas", "hares", "katib"],
        "priority": 5,
    },
    
    # Attendance
    "attendance": {
        "patterns": ["attendance", "حضور", "غياب", "check"],
        "agents": ["fahas", "katib"],
        "priority": 6,
    },
    
    # Overtime
    "overtime": {
        "patterns": ["overtime", "إضافي", "ساعات"],
        "agents": ["fahas", "mudaqqiq", "katib"],
        "priority": 7,
    },
    
    # Loans
    "loans": {
        "patterns": ["loan", "قرض", "سلف", "installment"],
        "agents": ["fahas", "hares", "katib"],
        "priority": 8,
    },
    
    # Disciplinary
    "disciplinary": {
        "patterns": ["disciplinary", "جزاء", "عقوبة", "تأديب"],
        "agents": ["fahas", "hares", "sijil", "katib"],
        "priority": 9,
    },
    
    # End of Service
    "end_of_service": {
        "patterns": ["end", "service", "نهاية", "خدمة", "مكافأة"],
        "agents": ["fahas", "mudaqqiq", "katib"],
        "priority": 10,
    },
    
    # Appraisals
    "appraisals": {
        "patterns": ["appraisal", "تقييم", "أداء", "performance"],
        "agents": ["fahas", "katib"],
        "priority": 11,
    },
    
    # Recruitment
    "recruitment": {
        "patterns": ["recruitment", "توظيف", "تعيين", "candidate"],
        "agents": ["fahas", "musleh", "katib"],
        "priority": 12,
    },
    
    # GR Transactions
    "gr_transactions": {
        "patterns": ["gr", "transaction", "معاملة"],
        "agents": ["fahas", "katib"],
        "priority": 13,
    },
    
    # Special sheets (لا تُعالج)
    "reference": {
        "patterns": ["reference", "مرجع", "lookup"],
        "agents": [],  # لا تُعالج - بيانات مرجعية فقط
        "priority": 999,
    },
    
    "instructions": {
        "patterns": ["instruction", "تعليمات", "guide", "help"],
        "agents": [],  # لا تُعالج - تعليمات فقط
        "priority": 999,
    },
}


class TemplateType:
    """نوع الـ Template"""
    MASTER = "master"  # ملف شامل متعدد الـ sheets
    MODULAR = "modular"  # ملف واحد لغرض واحد
    UNKNOWN = "unknown"


class SheetInfo:
    """معلومات عن sheet واحد"""
    def __init__(
        self,
        name: str,
        sheet_type: str,
        row_count: int,
        agents: List[str],
        priority: int,
    ):
        self.name = name
        self.sheet_type = sheet_type
        self.row_count = row_count
        self.agents = agents
        self.priority = priority
    
    def to_dict(self) -> Dict[str, Any]:
        return {
            "name": self.name,
            "sheet_type": self.sheet_type,
            "row_count": self.row_count,
            "agents": self.agents,
            "priority": self.priority,
        }


class TemplateDetectionResult:
    """نتيجة فحص الـ Template"""
    def __init__(
        self,
        file_path: str,
        template_type: str,
        sheets: List[SheetInfo],
        total_rows: int,
        requires_agents: Set[str],
    ):
        self.file_path = file_path
        self.template_type = template_type
        self.sheets = sheets
        self.total_rows = total_rows
        self.requires_agents = requires_agents
    
    def to_dict(self) -> Dict[str, Any]:
        return {
            "file_path": self.file_path,
            "template_type": self.template_type,
            "sheets": [s.to_dict() for s in self.sheets],
            "total_rows": self.total_rows,
            "requires_agents": sorted(list(self.requires_agents)),
        }


class TemplateDetector:
    """
    كاشف نوع الـ Template
    
    يفحص أي ملف Excel ويحدد:
    - Master أم Modular
    - أي sheets موجودة
    - أي وكلاء مطلوبين
    """
    
    def __init__(self):
        self.sheet_type_map = SHEET_TYPE_MAP
    
    def detect(self, file_path: str) -> TemplateDetectionResult:
        """
        فحص ملف Excel وتحديد نوعه
        
        Returns:
            TemplateDetectionResult مع كل المعلومات
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        # تحميل الـ workbook (read-only للسرعة)
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            logger.error(f"Failed to load workbook: {e}")
            raise
        
        sheet_names = wb.sheetnames
        logger.info(f"Detected {len(sheet_names)} sheets: {sheet_names}")
        
        # فحص كل sheet
        sheets: List[SheetInfo] = []
        total_rows = 0
        requires_agents: Set[str] = set()
        
        for sheet_name in sheet_names:
            ws = wb[sheet_name]
            row_count = ws.max_row - 1  # -1 للـ header
            
            # تحديد نوع الـ sheet
            sheet_type, agents, priority = self._identify_sheet_type(sheet_name)
            
            if sheet_type and agents:  # لو sheet مُعرّف
                sheet_info = SheetInfo(
                    name=sheet_name,
                    sheet_type=sheet_type,
                    row_count=row_count,
                    agents=agents,
                    priority=priority,
                )
                sheets.append(sheet_info)
                total_rows += row_count
                requires_agents.update(agents)
                
                logger.info(
                    f"Sheet '{sheet_name}': type={sheet_type}, "
                    f"rows={row_count}, agents={agents}"
                )
        
        # تحديد نوع الـ Template
        template_type = self._determine_template_type(sheets)
        
        # ترتيب الـ sheets حسب الأولوية
        sheets.sort(key=lambda s: s.priority)
        
        result = TemplateDetectionResult(
            file_path=str(path),
            template_type=template_type,
            sheets=sheets,
            total_rows=total_rows,
            requires_agents=requires_agents,
        )
        
        logger.info(
            f"Template detected: type={template_type}, "
            f"sheets={len(sheets)}, rows={total_rows}, "
            f"agents={sorted(requires_agents)}"
        )
        
        return result
    
    def _identify_sheet_type(
        self,
        sheet_name: str
    ) -> tuple[Optional[str], List[str], int]:
        """
        تحديد نوع الـ sheet بناءً على اسمه
        
        Returns:
            (sheet_type, agents, priority)
        """
        sheet_name_lower = sheet_name.lower()
        
        # تجربة كل الأنواع
        for sheet_type, config in self.sheet_type_map.items():
            patterns = config["patterns"]
            
            # تحقق من وجود أي pattern
            for pattern in patterns:
                if pattern.lower() in sheet_name_lower:
                    return (
                        sheet_type,
                        config["agents"],
                        config["priority"]
                    )
        
        # لو ما لقينا تطابق
        logger.warning(f"Unknown sheet type for: {sheet_name}")
        return None, [], 999
    
    def _determine_template_type(self, sheets: List[SheetInfo]) -> str:
        """
        تحديد نوع الـ Template: Master أم Modular
        
        Master = أكثر من sheet للمعالجة
        Modular = sheet واحد فقط
        """
        processable_sheets = [s for s in sheets if s.agents]
        
        if len(processable_sheets) == 0:
            return TemplateType.UNKNOWN
        elif len(processable_sheets) == 1:
            return TemplateType.MODULAR
        else:
            return TemplateType.MASTER


# ═══════════════════════════════════════════════════════════════
# Helper Functions
# ═══════════════════════════════════════════════════════════════

def detect_template(file_path: str) -> TemplateDetectionResult:
    """
    Quick helper لفحص template
    
    Usage:
        result = detect_template("file.xlsx")
        print(f"Type: {result.template_type}")
        print(f"Agents needed: {result.requires_agents}")
    """
    detector = TemplateDetector()
    return detector.detect(file_path)


def get_processing_plan(
    detection_result: TemplateDetectionResult
) -> Dict[str, Any]:
    """
    تحويل نتيجة الكشف لخطة معالجة
    
    Returns:
        خطة جاهزة للتنفيذ مع ترتيب الـ sheets والوكلاء
    """
    plan = {
        "template_type": detection_result.template_type,
        "total_rows": detection_result.total_rows,
        "processing_order": [],
    }
    
    for sheet in detection_result.sheets:
        if sheet.agents:  # فقط الـ sheets اللي تحتاج معالجة
            plan["processing_order"].append({
                "sheet_name": sheet.name,
                "sheet_type": sheet.sheet_type,
                "row_count": sheet.row_count,
                "agents": sheet.agents,
                "priority": sheet.priority,
            })
    
    return plan


# ═══════════════════════════════════════════════════════════════
# Testing
# ═══════════════════════════════════════════════════════════════

if __name__ == "__main__":
    import sys
    
    logging.basicConfig(level=logging.INFO)
    
    if len(sys.argv) < 2:
        print("Usage: python template_detector.py <file.xlsx>")
        sys.exit(1)
    
    file_path = sys.argv[1]
    
    try:
        result = detect_template(file_path)
        
        print("\n" + "="*60)
        print("TEMPLATE DETECTION RESULT")
        print("="*60)
        
        print(f"\nFile: {result.file_path}")
        print(f"Type: {result.template_type}")
        print(f"Total Rows: {result.total_rows:,}")
        print(f"Agents Needed: {sorted(result.requires_agents)}")
        
        print(f"\nSheets ({len(result.sheets)}):")
        for sheet in result.sheets:
            print(f"  [{sheet.priority}] {sheet.name}")
            print(f"      Type: {sheet.sheet_type}")
            print(f"      Rows: {sheet.row_count:,}")
            print(f"      Agents: {sheet.agents}")
        
        print("\n" + "="*60)
        print("PROCESSING PLAN")
        print("="*60)
        
        plan = get_processing_plan(result)
        
        for item in plan["processing_order"]:
            print(f"\n{item['priority']}. {item['sheet_name']}")
            print(f"   Type: {item['sheet_type']}")
            print(f"   Rows: {item['row_count']:,}")
            print(f"   Agents: {' → '.join(item['agents'])}")
        
        print("\n" + "="*60)
        
    except Exception as e:
        print(f"ERROR: {e}")
        sys.exit(1)
