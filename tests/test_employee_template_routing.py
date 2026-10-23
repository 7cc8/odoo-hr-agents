"""The official Employees template has an "Instructions" sheet, which used to send it to the
lenient excel_template_validator (company-only check, different row_hash). It now goes through
the strict pipeline (parser + Validator) that writer.py re-computes; other HR domains keep the
template validator."""
import openpyxl
import pytest

import dry_run_report
import generate_employee_template as gen
from parser import (
    OFFICIAL_TEMPLATE_EXAMPLE_EMAIL, OFFICIAL_TEMPLATE_EXAMPLE_NAME, parse_xlsx,
)
from tests.test_validator import FakeRepository

COMPANY = "شركة النخبة الذكية للتقنية"
COMPANY_REF = [15, COMPANY]
COMPANY_REC = {"id": 15, "name": COMPANY, "active": True}
DEPT = {"id": 8, "name": "الموارد البشرية", "complete_name": "الإدارة العليا / الموارد البشرية",
        "company_id": COMPANY_REF, "active": True}
JOB = {"id": 22, "name": "أخصائي موارد بشرية", "company_id": COMPANY_REF, "department_id": False, "active": True}
EXISTING = {"id": 91, "name": "موظف قديم", "work_email": "dup@example.com", "company_id": COMPANY_REF, "active": True}
SNAPSHOT = {"companies": [COMPANY_REC], "departments": [DEPT], "jobs": [JOB]}


class Repo(FakeRepository):
    def get_session_info(self):
        return {"odoo_url": "http://odoo.test", "odoo_db": "test", "uid": 1, "server_version": "19"}

    def get_company_ids(self):
        return [15]


@pytest.fixture
def repo(monkeypatch):
    r = Repo([COMPANY_REC], [DEPT], [JOB], [EXISTING])
    monkeypatch.setattr(dry_run_report, "OdooReadonlyRepository", lambda *a, **k: r)
    return r


def _row(name, email, mobile, dept="الإدارة العليا / الموارد البشرية", job="أخصائي موارد بشرية"):
    return [name, email, "", mobile, "محاسب", "2026-02-01", COMPANY, dept, job]


def _filled_template(tmp_path, rows, keep_example=True):
    path = gen.generate_template(SNAPSHOT, str(tmp_path))
    wb = openpyxl.load_workbook(path)
    ws = wb["Employees"]
    if not keep_example:
        for c in ws[2]:
            c.value = None
    for i, values in enumerate(rows, start=3):
        for j, v in enumerate(values, start=1):
            ws.cell(row=i, column=j, value=v)
    wb.save(path)
    return path


ROWS = [
    _row("سارة الحربي", "sara@example.com", "0500000001"),
    _row("خالد", "not-an-email", "0500000002", dept="قسم غير موجود", job="وظيفة غير موجودة"),
    _row("نورة", "dup@example.com", "0500000003"),
]


def test_official_template_uses_the_strict_pipeline(tmp_path, repo):
    path = _filled_template(tmp_path, ROWS)
    assert "Instructions" in openpyxl.load_workbook(path).sheetnames  # the reason it used to be diverted
    report = dry_run_report.build_full_report(path)
    by_row = {r["excel_row_number"]: r for r in report["rows"]}
    assert by_row[3]["final_status"] == "accepted"
    assert (by_row[4]["final_status"], by_row[4]["failure_code"]) == ("rejected", "invalid_work_email")
    assert (by_row[5]["final_status"], by_row[5]["failure_code"]) == ("needs_review", "duplicate_existing_employee")


def test_report_hashes_are_the_ones_writer_py_recomputes(tmp_path, repo):
    path = _filled_template(tmp_path, ROWS)
    report = dry_run_report.build_full_report(path)
    parsed_hashes = {r["row_hash"] for r in parse_xlsx(path)["rows"]}
    assert report["rows"], "report has no rows"
    assert {r["row_hash"] for r in report["rows"]} <= parsed_hashes
    assert report["header"]["source_hash"] == parse_xlsx(path)["source_hash"]


def test_untouched_example_row_is_skipped(tmp_path, repo):
    path = _filled_template(tmp_path, ROWS, keep_example=True)
    parsed = parse_xlsx(path)
    assert [r["excel_row_number"] for r in parsed["rows"]] == [3, 4, 5]
    assert "template_example_row_skipped:2" in parsed["warnings"]
    assert 2 not in {r["excel_row_number"] for r in dry_run_report.build_full_report(path)["rows"]}


def test_edited_example_row_is_real_data(tmp_path, repo):
    path = _filled_template(tmp_path, ROWS)
    ws = openpyxl.load_workbook(path)["Employees"]
    wb = ws.parent
    ws["B2"].value = "ahmed.new@example.com"  # same name, different e-mail -> a real employee
    wb.save(path)
    parsed = parse_xlsx(path)
    assert 2 in {r["excel_row_number"] for r in parsed["rows"]}
    assert "template_example_row_skipped:2" not in parsed["warnings"]


def test_employee_with_example_name_but_other_email_is_not_dropped(tmp_path, repo):
    path = _filled_template(tmp_path, ROWS, keep_example=False)
    wb = openpyxl.load_workbook(path)
    ws = wb["Employees"]
    ws["A2"].value = OFFICIAL_TEMPLATE_EXAMPLE_NAME
    ws["B2"].value = "different@example.com"
    ws["G2"].value, ws["H2"].value, ws["I2"].value = COMPANY, DEPT["complete_name"], JOB["name"]
    wb.save(path)
    assert 2 in {r["excel_row_number"] for r in parse_xlsx(path)["rows"]}


def test_example_constants_match_the_generator(tmp_path):
    path = gen.generate_template(SNAPSHOT, str(tmp_path))
    ws = openpyxl.load_workbook(path)["Employees"]
    assert ws["A2"].value == OFFICIAL_TEMPLATE_EXAMPLE_NAME
    assert ws["B2"].value == OFFICIAL_TEMPLATE_EXAMPLE_EMAIL


def _other_domain_workbook(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Time_Off"
    ws.append(["employee_id", "leave_type", "date_from", "date_to"])
    ws.append(["1088776655", "annual", "2026-03-01", "2026-03-05"])
    wb.create_sheet("Instructions")
    path = tmp_path / "leave.xlsx"
    wb.save(path)
    return path


def test_other_domains_keep_the_template_validator(tmp_path, repo, monkeypatch):
    import excel_template_validator as etv

    path = _other_domain_workbook(tmp_path)
    assert dry_run_report._detect_template_domain(path) != "employee"
    monkeypatch.setattr(etv, "validate_template_workbook", lambda *a, **k: {"routed": "template-validator"})
    monkeypatch.setattr(dry_run_report, "parse_xlsx", lambda *a, **k: pytest.fail("strict pipeline must not run"))
    assert dry_run_report.build_full_report(path) == {"routed": "template-validator"}


def test_structural_parser_error_falls_back_to_template_validator(tmp_path, repo, monkeypatch):
    import excel_template_validator as etv

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(["something", "else"])  # no employee headers -> ParserError
    wb.create_sheet("Instructions")
    path = tmp_path / "odd.xlsx"
    wb.save(path)
    monkeypatch.setattr(etv, "validate_template_workbook", lambda *a, **k: {"routed": "fallback"})
    assert dry_run_report.build_full_report(path) == {"routed": "fallback"}


def test_a_bug_in_the_strict_pipeline_does_not_silently_downgrade(tmp_path, repo, monkeypatch):
    import excel_template_validator as etv

    path = _filled_template(tmp_path, ROWS)
    monkeypatch.setattr(dry_run_report.Validator, "validate_rows", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
    called = []
    monkeypatch.setattr(etv, "validate_template_workbook", lambda *a, **k: called.append(1) or {})
    with pytest.raises(RuntimeError, match="boom"):
        dry_run_report.build_full_report(path)
    assert not called
