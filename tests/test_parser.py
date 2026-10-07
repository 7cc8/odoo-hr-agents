import unittest
import tempfile
import os
from openpyxl import Workbook
from parser import parse_xlsx, ParserError

HEADERS = ["name", "work_email", "work_phone", "mobile_phone", "job_title", "company", "department", "job"]


def build_xlsx(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in HEADERS])
    wb.save(path)


class ParserTests(unittest.TestCase):

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def _path(self, name):
        return os.path.join(self.tmpdir, name)

    def test_point5_possible_duplicate_in_file_name_case_and_phone_format(self):
        path = self._path("t1.xlsx")
        build_xlsx([
            {"name": "Ahmed", "work_email": "", "mobile_phone": "0500000000", "company": "شركة النخبة الذكية للتقنية"},
            {"name": "ahmed", "work_email": "", "mobile_phone": "+966 50 000 0000", "company": "شركة النخبة الذكية للتقنية"},
        ], path)
        report = parse_xlsx(path)
        flags = [set(r["parser_flags"]) for r in report["rows"]]
        self.assertTrue(all("possible_duplicate_in_file" in f for f in flags))

    def test_duplicate_in_file_by_email(self):
        path = self._path("t2.xlsx")
        build_xlsx([
            {"name": "أحمد", "work_email": "ahmed@example.com", "mobile_phone": "0500000001", "company": "شركة النخبة الذكية للتقنية"},
            {"name": "أحمد الثاني", "work_email": "Ahmed@Example.com ", "mobile_phone": "0500000002", "company": "شركة النخبة الذكية للتقنية"},
        ], path)
        report = parse_xlsx(path)
        flags = [set(r["parser_flags"]) for r in report["rows"]]
        self.assertTrue(all("duplicate_in_file" in f for f in flags))

    def test_blank_email_two_rows_not_duplicate_key(self):
        path = self._path("t3.xlsx")
        build_xlsx([
            {"name": "سالم", "work_email": "", "mobile_phone": "0500000003", "company": "شركة النخبة الذكية للتقنية"},
            {"name": "خالد", "work_email": "", "mobile_phone": "0500000004", "company": "شركة النخبة الذكية للتقنية"},
        ], path)
        report = parse_xlsx(path)
        for r in report["rows"]:
            self.assertIsNone(r["duplicate_key"])

    def test_missing_required_column_raises(self):
        path = self._path("t4.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Employees"
        ws.append(["name", "company"])
        ws.append(["أحمد", "شركة النخبة الذكية للتقنية"])
        wb.save(path)
        with self.assertRaises(ParserError):
            parse_xlsx(path)

    def test_hiring_date_parsed_when_column_present(self):
        path = self._path("t_hiring_date.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Employees"
        headers_with_date = HEADERS + ["hiring_date"]
        ws.append(headers_with_date)
        ws.append(["سارة", "sara@example.com", "", "0501112233", "مدير", "شركة النخبة", "IT", "مهندس", "2026-10-06"])
        wb.save(path)
        report = parse_xlsx(path)
        self.assertEqual(len(report["rows"]), 1)
        fields = report["rows"][0]["fields"]
        self.assertIn("hiring_date", fields)
        self.assertEqual(fields["hiring_date"]["value"], "2026-10-06")

    def test_hiring_date_empty_when_column_absent(self):
        path = self._path("t_no_hiring_date.xlsx")
        build_xlsx([
            {"name": "سارة", "work_email": "sara2@example.com", "mobile_phone": "0501112234", "company": "شركة النخبة"},
        ], path)
        report = parse_xlsx(path)
        self.assertEqual(len(report["rows"]), 1)
        fields = report["rows"][0]["fields"]
        self.assertIn("hiring_date", fields)
        self.assertEqual(fields["hiring_date"]["value"], "")

    def test_point6_run_id_source_hash_row_hash_present(self):
        path = self._path("t5.xlsx")
        build_xlsx([
            {"name": "أحمد", "work_email": "ahmed2@example.com", "mobile_phone": "0500000005", "company": "شركة النخبة الذكية للتقنية"},
        ], path)
        report = parse_xlsx(path)
        self.assertTrue(report["run_id"])
        self.assertTrue(report["source_hash"])
        self.assertTrue(report["rows"][0]["row_hash"])


if __name__ == "__main__":
    unittest.main()