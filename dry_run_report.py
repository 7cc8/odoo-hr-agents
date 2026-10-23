from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Dict, List, Optional
from zipfile import ZIP_DEFLATED, ZipFile
from xml.sax.saxutils import escape

from parser import ParserError, parse_xlsx
from validator import Validator
from odoo_readonly_repository import OdooReadonlyRepository


REPORT_HEADERS = [
    "excel_row_number",
    "final_status",
    "failure_code",
    "reasons",
    "company_id",
    "department_id",
    "job_id",
    "employee_match_id",
    "row_hash",
    "run_id",
    "source_hash",
    "parser_warnings",
]

CSV_INJECTION_TRIGGER_CHARS = ("=", "+", "-", "@")


class DryRunReportError(Exception):
    pass


def csv_formula_guard(value: str) -> str:
    """Prefix with a leading apostrophe when the CSV cell text starts with
    a formula/DDE trigger character. This is the standard CSV-injection
    mitigation; the apostrophe is intentionally visible/kept in CSV output.
    """
    if value and value[0] in CSV_INJECTION_TRIGGER_CHARS:
        return "'" + value
    return value


class _MiniXlsxEmitter:
    def __init__(self, file_path: str | Path):
        self.file_path = Path(file_path)

    def emit(self, rows: List[List[Any]], sheet_name: str = "DryRun") -> None:
        self.file_path.parent.mkdir(parents=True, exist_ok=True)

        content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
  <Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
  <Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
</Types>"""

        rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>"""

        workbook = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
          xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets>
    <sheet name="{escape(sheet_name)}" sheetId="1" r:id="rId1"/>
  </sheets>
</workbook>"""

        workbook_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>"""

        styles = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <fonts count="1">
    <font><sz val="11"/><name val="Calibri"/></font>
  </fonts>
  <fills count="1">
    <fill><patternFill patternType="none"/></fill>
  </fills>
  <borders count="1">
    <border/>
  </borders>
  <cellStyleXfs count="1">
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0"/>
  </cellStyleXfs>
  <cellXfs count="1">
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
  </cellXfs>
  <cellStyles count="1">
    <cellStyle name="Normal" xfId="0" builtinId="0"/>
  </cellStyles>
</styleSheet>"""

        core = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
                   xmlns:dc="http://purl.org/dc/elements/1.1/"
                   xmlns:dcterms="http://purl.org/dc/terms/"
                   xmlns:dcmitype="http://purl.org/dc/dcmitype/"
                   xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <dc:creator>dry_run_report</dc:creator>
  <cp:lastModifiedBy>dry_run_report</cp:lastModifiedBy>
</cp:coreProperties>"""

        app = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
            xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">
  <Application>Python</Application>
</Properties>"""

        def col_name(n: int) -> str:
            out = ""
            while n > 0:
                n, rem = divmod(n - 1, 26)
                out = chr(65 + rem) + out
            return out

        sheet_rows = []
        for r_idx, row in enumerate(rows, start=1):
            cells = []
            for c_idx, value in enumerate(row, start=1):
                ref = f"{col_name(c_idx)}{r_idx}"
                # inlineStr type is explicitly a string type: Excel will
                # NEVER evaluate this content as a formula, so no leading
                # apostrophe is needed or wanted here.
                text = "" if value is None else str(value)
                cells.append(
                    f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{escape(text)}</t></is></c>'
                )
            sheet_rows.append(f'<row r="{r_idx}">' + "".join(cells) + "</row>")

        sheet_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <sheetData>
    {rows_xml}
  </sheetData>
</worksheet>""".format(rows_xml="".join(sheet_rows))

        with ZipFile(self.file_path, "w", compression=ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", content_types)
            zf.writestr("_rels/.rels", rels)
            zf.writestr("docProps/core.xml", core)
            zf.writestr("docProps/app.xml", app)
            zf.writestr("xl/workbook.xml", workbook)
            zf.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
            zf.writestr("xl/styles.xml", styles)
            zf.writestr("xl/worksheets/sheet1.xml", sheet_xml)


def resolved_id(refs: Optional[Dict[str, Any]], key: str) -> Optional[int]:
    if not refs:
        return None
    item = refs.get(key)
    if isinstance(item, dict):
        value = item.get("id")
        if isinstance(value, int):
            return value
    return None


def stringify_reasons(reasons: Any) -> str:
    if reasons is None:
        return ""
    if isinstance(reasons, list):
        return " | ".join(str(x) for x in reasons)
    return str(reasons)


def stringify_parser_warnings(parser_flags: Any) -> str:
    if parser_flags is None:
        return ""
    if isinstance(parser_flags, list):
        return " | ".join(str(x) for x in parser_flags)
    return str(parser_flags)


def build_report_dict(
    parsed_report: Dict[str, Any],
    validated_rows: List[Dict[str, Any]],
    repo_info: Dict[str, Any],
    snapshot: Dict[str, Any],
) -> Dict[str, Any]:
    status_counts: Dict[str, int] = {}
    report_rows = []

    parser_row_map = {row["excel_row_number"]: row for row in parsed_report["rows"]}

    for validated in validated_rows:
        status = validated.get("final_status")
        status_counts[status] = status_counts.get(status, 0) + 1
        parsed_row = parser_row_map.get(validated["excel_row_number"], {})
        refs = validated.get("resolved_references") or {}

        report_rows.append(
            {
                "excel_row_number": validated.get("excel_row_number"),
                "final_status": status,
                "failure_code": validated.get("failure_code"),
                "reasons": validated.get("reasons") or [],
                "resolved_ids": {
                    "company_id": resolved_id(refs, "company"),
                    "department_id": resolved_id(refs, "department"),
                    "job_id": resolved_id(refs, "job"),
                    "employee_match_id": resolved_id(refs, "employee_match"),
                },
                "row_hash": validated.get("row_hash"),
                "run_id": validated.get("run_id"),
                "source_hash": validated.get("source_hash"),
                "parser_warnings": parsed_row.get("parser_flags") or [],
            }
        )

    header = {
        "odoo_url": repo_info.get("odoo_url"),
        "odoo_db": repo_info.get("odoo_db"),
        "uid": repo_info.get("uid"),
        "version": repo_info.get("server_version"),
        "reference_counts": {
            "companies": len(snapshot.get("companies", [])),
            "departments": len(snapshot.get("departments", [])),
            "jobs": len(snapshot.get("jobs", [])),
        },
        "run_id": parsed_report.get("run_id"),
        "source_hash": parsed_report.get("source_hash"),
    }

    return {
        "header": header,
        "summary": {
            "status_counts": status_counts,
            "row_count": len(report_rows),
        },
        "rows": report_rows,
    }


def write_csv(report_dict: Dict[str, Any], file_path: str | Path) -> None:
    file_path = Path(file_path)
    file_path.parent.mkdir(parents=True, exist_ok=True)

    with file_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=REPORT_HEADERS)
        writer.writeheader()

        for row in report_dict["rows"]:
            raw = {
                "excel_row_number": row.get("excel_row_number"),
                "final_status": row.get("final_status"),
                "failure_code": row.get("failure_code") or "",
                "reasons": stringify_reasons(row.get("reasons")),
                "company_id": row.get("resolved_ids", {}).get("company_id") or "",
                "department_id": row.get("resolved_ids", {}).get("department_id") or "",
                "job_id": row.get("resolved_ids", {}).get("job_id") or "",
                "employee_match_id": row.get("resolved_ids", {}).get("employee_match_id") or "",
                "row_hash": row.get("row_hash") or "",
                "run_id": row.get("run_id") or "",
                "source_hash": row.get("source_hash") or "",
                "parser_warnings": stringify_parser_warnings(row.get("parser_warnings")),
            }
            guarded = {k: csv_formula_guard(str(v)) for k, v in raw.items()}
            writer.writerow(guarded)


def write_json(report_dict: Dict[str, Any], file_path: str | Path) -> None:
    file_path = Path(file_path)
    file_path.parent.mkdir(parents=True, exist_ok=True)
    file_path.write_text(json.dumps(report_dict, ensure_ascii=False, indent=2), encoding="utf-8")


def write_xlsx(report_dict: Dict[str, Any], file_path: str | Path) -> None:
    rows = [REPORT_HEADERS]
    for row in report_dict["rows"]:
        rows.append(
            [
                row.get("excel_row_number"),
                row.get("final_status"),
                row.get("failure_code") or "",
                stringify_reasons(row.get("reasons")),
                row.get("resolved_ids", {}).get("company_id") or "",
                row.get("resolved_ids", {}).get("department_id") or "",
                row.get("resolved_ids", {}).get("job_id") or "",
                row.get("resolved_ids", {}).get("employee_match_id") or "",
                row.get("row_hash") or "",
                row.get("run_id") or "",
                row.get("source_hash") or "",
                stringify_parser_warnings(row.get("parser_warnings")),
            ]
        )
    _MiniXlsxEmitter(file_path).emit(rows)


def _detect_template_domain(input_xlsx_path: str | Path) -> str | None:
    """Domain key ("employee", "contract", ...) of a system template, or None if undetectable.

    None sends the file to the template validator, which is the behaviour before routing existed.
    """
    try:
        import openpyxl
        from excel_template_validator import identify_template_and_sheet

        wb = openpyxl.load_workbook(str(input_xlsx_path), data_only=True)
        return identify_template_and_sheet(wb, Path(input_xlsx_path).name)[0]
    except Exception:
        return None


def build_full_report(input_xlsx_path: str | Path) -> Dict[str, Any]:
    """Parses, validates, and assembles the report dict. Does NOT write
    any output files and does NOT call create/write/unlink on Odoo.
    Supports both legacy employee import files and all 14 modular/master HR templates."""
    repository = OdooReadonlyRepository()

    # 1. Check if the file is one of the modular/master HR templates
    try:
        from excel_template_validator import validate_template_workbook
        import openpyxl
        wb = openpyxl.load_workbook(str(input_xlsx_path), read_only=True)
        sheetnames = wb.sheetnames
        is_template = (
            "Instructions" in sheetnames
            or any(
                s in sheetnames
                for s in [
                    "Attendances", "Contracts", "Payroll_GOSI", "Time_Off",
                    "Overtime", "Loans", "Disciplinary", "End_of_Service",
                    "Appraisals", "Recruitment", "GR_Transactions", "Company_Setup"
                ]
            )
            or any(s.startswith(("1_", "2_", "3_", "4_", "5_", "6_", "7_", "8_", "9_", "10_", "11_", "12_", "13_")) for s in sheetnames)
            or any(k in Path(input_xlsx_path).name.lower() for k in ["template_", "master_", "attendance", "contract", "payroll", "timeoff", "loan", "overtime", "disciplinary", "appraisal", "recruitment"])
        )
        if is_template and _detect_template_domain(input_xlsx_path) != "employee":
            # Other HR domains (contracts, leave, payroll, ...) keep the template validator.
            # Employee templates fall through to the strict pipeline below.
            return validate_template_workbook(str(input_xlsx_path), repo=repository)
    except Exception:
        pass

    # 2. Strict employee pipeline (parser + Validator). Its row_hash is what writer.py re-computes.
    try:
        parsed_report = parse_xlsx(str(input_xlsx_path))
    except ParserError:
        # 3. Not employee-shaped (e.g. a template with renamed sheets): let the template validator try.
        # Only a structural ParserError falls back; any other failure (Odoo, a bug) must surface
        # instead of silently downgrading to the lenient template validator.
        from excel_template_validator import validate_template_workbook
        return validate_template_workbook(str(input_xlsx_path), repo=repository)
    snapshot = repository.get_live_reference_snapshot()
    validator = Validator(repository)
    validated_rows = validator.validate_rows(parsed_report["rows"])
    repo_info = repository.get_session_info()
    return build_report_dict(parsed_report, validated_rows, repo_info, snapshot)


def write_report_files(report_dict: Dict[str, Any], output_dir: str | Path) -> Dict[str, Path]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    paths = {
        "json": output_dir / "dry_run_report.json",
        "csv": output_dir / "dry_run_report.csv",
        "xlsx": output_dir / "dry_run_report.xlsx",
    }

    write_json(report_dict, paths["json"])
    write_csv(report_dict, paths["csv"])
    write_xlsx(report_dict, paths["xlsx"])

    return paths


def generate_dry_run_report(input_xlsx_path: str | Path, output_dir: str | Path) -> Dict[str, Any]:
    report_dict = build_full_report(input_xlsx_path)
    write_report_files(report_dict, output_dir)
    return report_dict