from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Dict, Any
from zipfile import ZipFile

from openpyxl import load_workbook

from normalization import (
    clean_value_text, to_text_preserve, normalize_header, normalize_date_text,
    email_match_key, name_match_key, reference_match_key, phone_match_key,
)

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB (كان 2 MB)
MAX_ROWS = 10000  # 10,000 صف (كان 1,000)
REQUIRED_SHEET = "Employees"  # fallback - الآن يدعم أي sheet
REQUIRED_HEADERS = ["name", "work_email", "work_phone", "mobile_phone", "job_title", "company", "department", "job"]
OPTIONAL_HEADERS = ["hiring_date"]
ALL_HEADERS = REQUIRED_HEADERS + OPTIONAL_HEADERS
PHONE_HEADERS = {"work_phone", "mobile_phone"}

# Values of the example row shipped in the modular Employees template
# (scripts/generate_all_hr_templates.py). Matched EXACTLY so a real employee
# whose name merely contains "عبدالله" is never skipped.
TEMPLATE_EXAMPLE_ID = "1088776655"
TEMPLATE_EXAMPLE_NAME = "عبدالله بن محمد الغامدي"

# Example row (row 2, header on row 1) of the official Employees template written by
# generate_employee_template.py. Skipped only when BOTH values match exactly, so an edited
# row (or a real employee with the same name but another e-mail) is never dropped.
OFFICIAL_TEMPLATE_EXAMPLE_NAME = "أحمد محمد العتيبي"
OFFICIAL_TEMPLATE_EXAMPLE_EMAIL = "ahmed@company.com"


class ParserError(Exception):
    pass


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def make_run_id(source_hash: str) -> str:
    return source_hash[:16]


def row_hash(normalized: Dict[str, Dict[str, Any]]) -> str:
    payload = json.dumps(normalized, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _match_key_for(header: str, value_text: str):
    if header == "work_email":
        return email_match_key(value_text)
    if header in PHONE_HEADERS:
        return phone_match_key(value_text)
    if header == "name":
        return name_match_key(value_text)
    return reference_match_key(value_text)


def _xlsx_has_formulas_or_merged_or_hidden(path: Path):
    warnings = []
    with ZipFile(path, "r") as zf:
        workbook_xml = zf.read("xl/workbook.xml").decode("utf-8", errors="ignore")
        if 'state="hidden"' in workbook_xml:
            warnings.append("hidden_sheet_present")
    wb = load_workbook(path, data_only=False)
    for ws in wb.worksheets:
        if ws.merged_cells.ranges:
            warnings.append(f"merged_cells:{ws.title}")
        hidden_rows = [idx for idx, dim in ws.row_dimensions.items() if getattr(dim, "hidden", False)]
        if hidden_rows:
            warnings.append(f"hidden_rows:{ws.title}:{','.join(map(str, hidden_rows[:20]))}")
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    warnings.append(f"formula_present:{ws.title}:{cell.coordinate}")
                    return sorted(set(warnings))
    return sorted(set(warnings))


def parse_xlsx(file_path: str) -> Dict[str, Any]:
    path = Path(file_path)
    if path.suffix.lower() != ".xlsx":
        raise ParserError("only_xlsx_supported")
    if not path.exists():
        raise ParserError("file_not_found")
    size = path.stat().st_size
    if size > MAX_FILE_SIZE_BYTES:
        raise ParserError("file_too_large")

    source_hash = sha256_bytes(path.read_bytes())
    run_id = make_run_id(source_hash)

    workbook = load_workbook(path, data_only=True)
    warnings = _xlsx_has_formulas_or_merged_or_hidden(path)

    target_sheet = REQUIRED_SHEET
    if target_sheet not in workbook.sheetnames:
        # Check for alternative data sheet
        aux = {"instructions", "reference", "دليل التعليمات", "المراجع"}
        candidates = [s for s in workbook.sheetnames if s.strip().lower() not in aux]
        emp_cand = [s for s in candidates if "employee" in s.lower() or "موظف" in s]
        if emp_cand:
            target_sheet = emp_cand[0]
        elif candidates:
            target_sheet = candidates[0]
        else:
            raise ParserError("missing_employees_sheet")

    ws = workbook[target_sheet]
    if ws.max_row - 1 > MAX_ROWS:
        raise ParserError("too_many_rows")

    header_row_idx = 1
    row1_non_empty = [c.value for c in ws[1] if c.value is not None]
    if len(row1_non_empty) <= 2 and ws.max_row >= 2:
        row2_non_empty = [c.value for c in ws[2] if c.value is not None]
        if len(row2_non_empty) >= 3:
            header_row_idx = 2

    header_cells = [cell.value for cell in ws[header_row_idx]]
    headers_norm = [normalize_header(to_text_preserve(v)) for v in header_cells]

    # Map Arabic headers if present
    AR_MAP = {
        "اسم الموظف الكامل بالعربي": "name",
        "اسم الموظف": "name",
        "البريد الإلكتروني للعمل": "work_email",
        "البريد الإلكتروني": "work_email",
        "رقم الجوال (05xxxxxxxx)": "mobile_phone",
        "رقم الجوال": "mobile_phone",
        "هاتف العمل": "work_phone",
        "المسمى الوظيفي": "job_title",
        "المنشأة التابع لها": "company",
        "الشركة": "company",
        "القسم / الإدارة": "department",
        "القسم": "department",
        "تاريخ المباشرة": "hiring_date",
    }
    for idx, h in enumerate(headers_norm):
        if not h:
            continue  # blank header cell: '' in <any str> is True and would be mislabeled 'name'
        for ar_k, en_k in AR_MAP.items():
            if ar_k in h or h in ar_k:
                headers_norm[idx] = en_k
                break

    if len(headers_norm) != len(set(headers_norm)):
        # Filter duplicates only if non-empty
        non_empty_headers = [h for h in headers_norm if h]
        if len(non_empty_headers) != len(set(non_empty_headers)):
            raise ParserError("duplicate_header")

    header_map = {h: idx + 1 for idx, h in enumerate(headers_norm) if h}
    if "job_title" in header_map and "job" not in header_map:
        header_map["job"] = header_map["job_title"]
    if "mobile_phone" in header_map and "work_phone" not in header_map:
        header_map["work_phone"] = header_map["mobile_phone"]

    if target_sheet == REQUIRED_SHEET or "employee" in target_sheet.lower():
        missing = [h for h in REQUIRED_HEADERS if h not in header_map]
        if missing:
            raise ParserError("missing_required_columns:" + ",".join(missing))

    rows = []
    duplicate_counts = {}
    possible_duplicate_counts = {}

    data_start_row = header_row_idx + 1
    if header_row_idx == 2 and ws.max_row > 3:
        # Check if row 3 is the template example row
        row3_vals = [to_text_preserve(cell.value).strip() for cell in ws[3] if cell.value is not None]
        if row3_vals and (TEMPLATE_EXAMPLE_ID in row3_vals or row3_vals[0] == TEMPLATE_EXAMPLE_NAME):
            data_start_row = 4
            warnings = sorted(set(warnings) | {"template_example_row_skipped:3"})
    elif header_row_idx == 1 and ws.max_row >= 2 and "name" in header_map and "work_email" in header_map:
        example_name = clean_value_text(to_text_preserve(ws.cell(row=2, column=header_map["name"]).value))
        example_email = clean_value_text(to_text_preserve(ws.cell(row=2, column=header_map["work_email"]).value)).lower()
        if example_name == OFFICIAL_TEMPLATE_EXAMPLE_NAME and example_email == OFFICIAL_TEMPLATE_EXAMPLE_EMAIL:
            data_start_row = 3
            warnings = sorted(set(warnings) | {"template_example_row_skipped:2"})

    for excel_row in range(data_start_row, ws.max_row + 1):
        row_raw, row_value, row_match = {}, {}, {}
        flags = []
        is_all_empty = True

        for header in ALL_HEADERS:
            if header in header_map:
                cell = ws.cell(row=excel_row, column=header_map[header])
                raw_text = to_text_preserve(cell.value)
                value_text = clean_value_text(raw_text)
                if header == "hiring_date":
                    value_text = normalize_date_text(value_text)
                match_text = _match_key_for(header, value_text)

                if header in PHONE_HEADERS and cell.data_type == "n" and raw_text:
                    flags.append("numeric_cell_in_text_column")
            else:
                raw_text = ""
                value_text = ""
                match_text = None

            row_raw[header] = raw_text
            row_value[header] = value_text
            row_match[header] = match_text
            if value_text != "" and header in REQUIRED_HEADERS:
                is_all_empty = False

        if is_all_empty:
            continue

        duplicate_key = row_match["work_email"] if row_match["work_email"] else None
        possible_duplicate_key = None
        if not duplicate_key and row_match["name"] and row_match["mobile_phone"]:
            possible_duplicate_key = row_match["name"] + "|" + row_match["mobile_phone"]

        if duplicate_key:
            duplicate_counts[duplicate_key] = duplicate_counts.get(duplicate_key, 0) + 1
        if possible_duplicate_key:
            possible_duplicate_counts[possible_duplicate_key] = possible_duplicate_counts.get(possible_duplicate_key, 0) + 1

        normalized_payload = {k: {"raw": row_raw[k], "value": row_value[k], "match_key": row_match[k]} for k in ALL_HEADERS}

        rows.append({
            "excel_row_number": excel_row,
            "run_id": run_id,
            "source_hash": source_hash,
            "row_hash": row_hash(normalized_payload),
            "fields": normalized_payload,
            "duplicate_key": duplicate_key,
            "possible_duplicate_key": possible_duplicate_key,
            "parser_flags": flags,
        })

    for row in rows:
        if row["duplicate_key"] and duplicate_counts.get(row["duplicate_key"], 0) > 1:
            row["parser_flags"].append("duplicate_in_file")
        if row["possible_duplicate_key"] and possible_duplicate_counts.get(row["possible_duplicate_key"], 0) > 1:
            row["parser_flags"].append("possible_duplicate_in_file")

    return {
        "run_id": run_id, "source_hash": source_hash, "file_size_bytes": size,
        "warnings": warnings, "row_count": len(rows), "rows": rows,
    }